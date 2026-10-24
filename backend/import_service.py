
import json
import csv
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
import zipfile
import os
import stat
from pathlib import PurePosixPath

import psycopg
from psycopg import sql


PROJECT_DIR = Path(__file__).resolve().parents[1]
PARSERS_DIR = PROJECT_DIR / "parsers"

# У каждого типа исходника свой готовый скрипт из папки parsers/.
# CSV намеренно не включён: у произвольного CSV нет единой структуры,
# поэтому нельзя безопасно угадывать, в какую таблицу его записывать.
PARSER_SCRIPTS = {
    ".xml": "PointzAgregator_parser.py",
    ".json": "parse_FrequentFlyerForum_Profiles_JSON.py",
    ".yaml": "SkyTeamYaml.py",
    ".yml": "SkyTeamYaml.py",
    ".tab": "parse_sirena_format_TAB.py",
    ".xlsx": "parse_boarding_passes.py",
    ".pdf": "parse_Skyteam_Timetable_PDF.py",
}

# Имена CSV, которые ожидают таблицы из data/sql/01_schema.sql.
STAGING_OUTPUTS = {
    ".xml": {
        "users.csv": "pointz_users.csv",
        "cards.csv": "pointz_cards.csv",
        "activities.csv": "pointz_activities.csv",
        "errors.csv": "pointz_errors.csv",
    },
    ".json": {
        "profiles.csv": "forum_profiles.csv",
        "documents.csv": "forum_documents.csv",
        "cards.csv": "forum_cards.csv",
        "flights.csv": "forum_flights.csv",
    },
    ".yaml": {
        "flights.csv": "skyteam_flights.csv",
        "flight_cards.csv": "skyteam_flight_cards.csv",
    },
    ".yml": {
        "flights.csv": "skyteam_flights.csv",
        "flight_cards.csv": "skyteam_flight_cards.csv",
    },
    ".tab": {"sirena_records.csv": "sirena_records.csv"},
    ".pdf": {"timetable_records.csv": "timetable_records.csv"},
}

BOARDING_DATA_COLUMNS = [
    "passengerfirstname", "passengersecondname", "passengerlastname",
    "passengersex", "passengerbirthdate", "passengerdocument",
    "bookingcode", "ticketnumber", "baggage", "flightdate",
    "flighttime", "flightnumber", "codeshare", "destination",
]

STAGING_TABLES = {
    "forum_profiles.csv": ("forum", "forum_profiles"),
    "forum_documents.csv": ("forum", "forum_documents"),
    "forum_cards.csv": ("forum", "forum_cards"),
    "forum_flights.csv": ("forum", "forum_flights"),
    "skyteam_flights.csv": ("skyteam", "skyteam_flights"),
    "skyteam_flight_cards.csv": ("skyteam", "skyteam_flight_cards"),
    "timetable_records.csv": ("skyteam", "timetable_records"),
    "sirena_records.csv": ("skyteam", "sirena_records"),
    "boarding_passes.csv": ("skyteam", "boarding_passes"),
    "boarding_data.csv": ("skyteam", "boarding_data"),
    "pointz_users.csv": ("pointz", "pointz_users"),
    "pointz_cards.csv": ("pointz", "pointz_cards"),
    "pointz_activities.csv": ("pointz", "pointz_activities"),
    "pointz_errors.csv": ("pointz", "pointz_errors"),
}

ZIP_MAX_ENTRIES = 5000
ZIP_MAX_FILE_SIZE = 5 * 1024**3
ZIP_MAX_TOTAL_SIZE = 20 * 1024**3


def _validate_zip(path: Path) -> None:
    """Проверяет архив, ограничения размера и безопасность путей."""
    if not zipfile.is_zipfile(path):
        raise ValueError(f"{path.name}: некорректный ZIP-архив")

    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()

            if len(entries) > ZIP_MAX_ENTRIES:
                raise ValueError(
                    f"{path.name}: слишком много элементов в архиве"
                )

            total_size = 0

            for entry in entries:
                name = entry.filename.replace("\\", "/")
                parts = PurePosixPath(name).parts

                if (
                    name.startswith("/")
                    or ".." in parts
                    or (parts and ":" in parts[0])
                ):
                    raise ValueError(
                        f"{path.name}: небезопасный путь в архиве: {name}"
                    )

                mode = (entry.external_attr >> 16) & 0xFFFF
                if stat.S_ISLNK(mode):
                    raise ValueError(
                        f"{path.name}: символические ссылки запрещены"
                    )

                if entry.file_size > ZIP_MAX_FILE_SIZE:
                    raise ValueError(
                        f"{name}: файл внутри архива слишком большой"
                    )

                total_size += entry.file_size

                if total_size > ZIP_MAX_TOTAL_SIZE:
                    raise ValueError(
                        f"{path.name}: распакованный архив слишком большой"
                    )

                if not entry.is_dir() and name.lower().endswith(".zip"):
                    raise ValueError(
                        "Вложенные ZIP-архивы пока не поддерживаются. "
                        "Распакуйте внутренний архив отдельно."
                    )

    except zipfile.BadZipFile as exc:
        raise ValueError(f"{path.name}: повреждённый ZIP-архив") from exc


def _extract_zip_safely(
    archive_path: Path,
    destination: Path,
) -> list[Path]:
    """Распаковывает архив в отдельную папку, не затрагивая исходник."""
    _validate_zip(archive_path)

    destination.mkdir(parents=True, exist_ok=True)
    root = destination.resolve()
    extracted = []

    with zipfile.ZipFile(archive_path) as archive:
        for entry in archive.infolist():
            if entry.is_dir():
                continue

            name = entry.filename.replace("\\", "/")
            parts = PurePosixPath(name).parts

            if not parts:
                continue

            target = (destination / Path(*parts)).resolve()

            if os.path.commonpath(
                [str(root), str(target)]
            ) != str(root):
                raise ValueError(
                    f"Небезопасный путь в ZIP: {name}"
                )

            target.parent.mkdir(parents=True, exist_ok=True)

            with archive.open(entry) as source:
                with target.open("wb") as output:
                    shutil.copyfileobj(
                        source,
                        output,
                        length=1024 * 1024,
                    )

            extracted.append(target)

    return extracted


def validate_file_content(path: Path, extension: str) -> None:
    """
    Проверяет, что содержимое файла хотя бы похоже
    на заявленный формат. Выбрасывает ValueError при ошибке.
    """
    extension = extension.lower()

    if extension == ".zip":
        _validate_zip(path)
        return

    with path.open("rb") as file:
        header = file.read(8192)

    if not header:
        raise ValueError(f"Файл {path.name} пустой")

    if extension == ".pdf":
        if not header.startswith(b"%PDF-"):
            raise ValueError(
                f"{path.name}: содержимое не похоже на PDF"
            )
        return

    if extension == ".xlsx":
        if not zipfile.is_zipfile(path):
            raise ValueError(
                f"{path.name}: это не корректный XLSX-архив"
            )

        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())

            if (
                "[Content_Types].xml" not in names
                or "xl/workbook.xml" not in names
            ):
                raise ValueError(
                    f"{path.name}: архив не похож на Excel XLSX"
                )
        return

    if extension in {".xml", ".json", ".yaml", ".yml",
                     ".tab", ".csv"}:
        if b"\x00" in header:
            raise ValueError(
                f"{path.name}: обнаружены нулевые байты; "
                "возможно, это не текстовый файл"
            )

        try:
            text = header.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(
                f"{path.name}: начало файла не читается как UTF-8"
            ) from exc

        if extension == ".xml":
            if "<" not in text:
                raise ValueError(
                    f"{path.name}: не найдено начало XML-разметки"
                )

        elif extension == ".json":
            first = text.lstrip()[:1]
            if first not in {"{", "["}:
                raise ValueError(
                    f"{path.name}: содержимое не похоже на JSON"
                )

        # Для YAML, CSV и TAB это предварительная проверка.
        # Полную корректность проверит соответствующий парсер.
        return

    raise ValueError(
        f"Проверка формата {extension} не реализована"
    )


def _parser_command(source: Path, output: Path) -> list[str]:
    """Строит команду конкретного готового парсера."""
    extension = source.suffix.lower()
    script_name = PARSER_SCRIPTS[extension]
    script = PARSERS_DIR / script_name

    if extension == ".xml":
        return [sys.executable, str(script), str(source), "--out", str(output)]

    return [
        sys.executable, str(script), "--input", str(source),
        "--output", str(output),
    ]


def _normalise_header(values: list) -> list[str]:
    """Приводит заголовки парсеров к именам колонок SQL."""
    return [
        str(value or "")
        .replace("\ufeff", "")
        .strip()
        .casefold()
        .replace("-", "_")
        for value in values
    ]


def _detect_csv_delimiter(path: Path) -> str:
    # Парсер посадочных талонов создаёт CSV с разделителем ;
    if path.name.endswith("_table.csv"):
        return ";"

    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(65536)

    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        return ","


def _max_int_column(path: Path, column: str) -> int:
    """Находит максимальный числовой ID в staging без загрузки всего файла."""
    if not path.is_file():
        return 0

    maximum = 0
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or column not in reader.fieldnames:
            return 0

        for row in reader:
            try:
                maximum = max(maximum, int(row.get(column) or 0))
            except ValueError:
                continue

    return maximum


def _append_csv(
    source: Path,
    target: Path,
    *,
    delimiter: str | None = None,
    expected_header: list[str] | None = None,
    id_offsets: dict[str, int] | None = None,
    value_prefixes: dict[str, str] | None = None,
) -> None:
    """Потоково добавляет CSV в staging, нормализуя заголовки.

    Сначала проверяет CSV, затем записывает строки порциями из чтения файла.
    Промежуточный список со всеми строками в памяти не создаётся.
    """
    delimiter = delimiter or _detect_csv_delimiter(source)

    with source.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        raw_header = next(reader, None)

        if not raw_header:
            raise ValueError(f"{source.name}: отсутствует заголовок CSV")

        header = _normalise_header(raw_header)

        if any(not column for column in header):
            raise ValueError(f"{source.name}: пустое имя колонки")

        if len(set(header)) != len(header):
            raise ValueError(
                f"{source.name}: повторяющиеся колонки после нормализации"
            )

        if expected_header is not None and header != expected_header:
            raise ValueError(
                f"{source.name}: ожидался заголовок {expected_header}, "
                f"получен {header}"
            )

        # Проверяем строки до записи, чтобы ошибочная строка
        # не оставила частично добавленные данные в staging.
        expected_width = len(raw_header)
        for line_number, row in enumerate(reader, start=2):
            if not row or (len(row) == 1 and not row[0].strip()):
                continue

            if len(row) != expected_width:
                raise ValueError(
                    f"{source.name}, строка {line_number}: ожидалось "
                    f"{expected_width} полей, получено {len(row)}"
                )

    target.parent.mkdir(parents=True, exist_ok=True)
    exists = target.is_file()

    if exists:
        with target.open("r", encoding="utf-8-sig", newline="") as stream:
            existing_header = _normalise_header(next(csv.reader(stream), []))

        if existing_header != header:
            raise ValueError(
                f"{source.name}: заголовок не совпадает с {target.name}"
            )

    offsets = id_offsets or {}
    prefixes = value_prefixes or {}
    indexes = {name: i for i, name in enumerate(header)}

    # Staging хранится с обычным CSV-разделителем — запятой.
    with source.open("r", encoding="utf-8-sig", newline="") as input_stream:
        reader = csv.reader(input_stream, delimiter=delimiter)
        next(reader, None)

        with target.open(
            "a" if exists else "w",
            encoding="utf-8",
            newline="",
        ) as output_stream:
            writer = csv.writer(output_stream, lineterminator="\n")

            if not exists:
                writer.writerow(header)

            for row in reader:
                if not row or (len(row) == 1 and not row[0].strip()):
                    continue

                for column, offset in offsets.items():
                    index = indexes.get(column)

                    if index is not None and row[index].strip() and offset:
                        try:
                            row[index] = str(int(row[index]) + offset)
                        except ValueError as exc:
                            raise ValueError(
                                f"{source.name}: нечисловой ID в колонке "
                                f"{column}: {row[index]!r}"
                            ) from exc

                for column, prefix in prefixes.items():
                    index = indexes.get(column)
                    if index is not None and row[index].strip():
                        row[index] = prefix + row[index]

                writer.writerow(row)

def _stage_parsed_files(
    source: Path,
    output: Path,
    staging: Path,
) -> list[str]:
    """Переносит результаты парсеров в staging для SQL-загрузки."""
    extension = source.suffix.lower()
    mapping = dict(STAGING_OUTPUTS.get(extension, {}))

    if extension == ".xlsx":
        mapping[f"{source.stem}_table.csv"] = "boarding_passes.csv"

    if not mapping:
        return []

    # Для каждого нового XML/YAML берём максимум уже записанных ID.
    pointz_offsets = {}
    yaml_flight_offset = 0

    if extension == ".xml":
        pointz_offsets = {
            "user_id": _max_int_column(
                staging / "pointz_users.csv", "user_id"
            ),
            "card_id": _max_int_column(
                staging / "pointz_cards.csv", "card_id"
            ),
            "activity_id": _max_int_column(
                staging / "pointz_activities.csv", "activity_id"
            ),
        }

    elif extension in {".yaml", ".yml"}:
        yaml_flight_offset = _max_int_column(
            staging / "skyteam_flights.csv", "flight_id"
        )

    staged = []

    for produced_name, staging_name in mapping.items():
        produced = output / produced_name

        if not produced.is_file():
            raise FileNotFoundError(
                f"{source.name}: парсер не создал ожидаемый файл "
                f"{produced_name}"
            )

        id_offsets = {}
        value_prefixes = {}

        if extension == ".xml":
            if staging_name == "pointz_users.csv":
                id_offsets = {
                    "user_id": pointz_offsets["user_id"]
                }

            elif staging_name == "pointz_cards.csv":
                id_offsets = {
                    "user_id": pointz_offsets["user_id"],
                    "card_id": pointz_offsets["card_id"],
                }

            elif staging_name == "pointz_activities.csv":
                id_offsets = {
                    "card_id": pointz_offsets["card_id"],
                    "activity_id": pointz_offsets["activity_id"],
                }

            elif staging_name == "pointz_errors.csv":
                id_offsets = {
                    "user_id": pointz_offsets["user_id"],
                    "card_id": pointz_offsets["card_id"],
                    "activity_id": pointz_offsets["activity_id"],
                }

        elif extension in {".yaml", ".yml"}:
            id_offsets = {
                "flight_id": yaml_flight_offset
            }

        elif extension == ".tab":
            # record_id должен оставаться различимым между исходниками.
            value_prefixes = {
                "record_id": f"{source.stem}:"
            }

        elif extension == ".pdf":
            value_prefixes = {
                "schedule_id": f"{source.stem}:"
            }

        target = staging / staging_name

        _append_csv(
            produced,
            target,
            id_offsets=id_offsets,
            value_prefixes=value_prefixes,
        )

        staged.append(f"staging/{staging_name}")

    return staged


def _stage_boarding_data_csv(
    source: Path,
    staging: Path,
) -> list[str]:
    """Распознаёт BoardingData CSV и добавляет его строки в общий staging."""
    delimiter = _detect_csv_delimiter(source)

    with source.open("r", encoding="utf-8-sig", newline="") as stream:
        header = _normalise_header(
            next(csv.reader(stream, delimiter=delimiter), [])
        )

    if header != BOARDING_DATA_COLUMNS:
        return []

    _append_csv(
        source,
        staging / "boarding_data.csv",
        delimiter=delimiter,
        expected_header=BOARDING_DATA_COLUMNS,
    )

    return ["staging/boarding_data.csv"]


def _stage_boarding_data_xlsx(
    source: Path,
    staging: Path,
) -> list[str]:
    """Распознаёт Excel-выгрузку BoardingData и добавляет все её строки."""
    from openpyxl import load_workbook

    target = staging / "boarding_data.csv"
    workbook = load_workbook(source, read_only=True, data_only=True)

    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            raw_header = next(rows, None)

            if raw_header is None:
                continue

            header = _normalise_header(list(raw_header))

            # Excel может хранить пустые колонки после последнего заголовка.
            while header and not header[-1]:
                header.pop()

            if header != BOARDING_DATA_COLUMNS:
                continue

            exists = target.is_file()

            if exists:
                with target.open(
                    "r", encoding="utf-8-sig", newline=""
                ) as stream:
                    existing_header = next(csv.reader(stream), [])

                if existing_header != BOARDING_DATA_COLUMNS:
                    raise ValueError(
                        "В staging/boarding_data.csv уже другой заголовок"
                    )

            with target.open(
                "a" if exists else "w",
                encoding="utf-8",
                newline="",
            ) as output:
                writer = csv.writer(output, lineterminator="\n")

                if not exists:
                    writer.writerow(BOARDING_DATA_COLUMNS)

                for row in rows:
                    if row is None:
                        continue

                    values = list(row[:len(BOARDING_DATA_COLUMNS)])
                    values.extend(
                        [""] * (len(BOARDING_DATA_COLUMNS) - len(values))
                    )

                    if any(
                        value is not None and str(value).strip()
                        for value in values
                    ):
                        writer.writerow(
                            "" if value is None else value
                            for value in values
                        )

            return ["staging/boarding_data.csv"]

    finally:
        workbook.close()

    return []


def _csv_header(path: Path) -> list[str]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        header = next(csv.reader(stream), None)
    if not header:
        raise ValueError(f"{path.name}: staging CSV пустой")
    return [column.strip() for column in header]


def _load_one_staging_file(
    connection: psycopg.Connection, path: Path, schema: str, table: str
) -> dict:
    """Загружает CSV через временную таблицу, не добавляя дубли."""
    columns = _csv_header(path)
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = %s AND table_name = %s
            """,
            (schema, table),
        )
        target_columns = {row[0] for row in cursor.fetchall()}
        unknown = [column for column in columns if column not in target_columns]
        if unknown:
            raise ValueError(
                f"{path.name}: неизвестные колонки для {schema}.{table}: {unknown}"
            )

        temp_name = f"import_rows_{schema}_{table}"
        target = sql.Identifier(schema, table)
        column_sql = sql.SQL(", ").join(map(sql.Identifier, columns))
        # В дампе могут быть служебные ID без DEFAULT. В staging их нет,
        # поэтому временная таблица содержит только реальные CSV-колонки.
        text_columns = sql.SQL(", ").join(
            sql.SQL("{} text").format(sql.Identifier(column))
            for column in columns
        )
        cursor.execute(
            sql.SQL("CREATE TEMP TABLE {} ({}) ON COMMIT DROP").format(
                sql.Identifier(temp_name), text_columns
            )
        )
        with cursor.copy(
            sql.SQL("COPY {} ({}) FROM STDIN WITH (FORMAT CSV, HEADER TRUE, ENCODING 'UTF8')")
            .format(sql.Identifier(temp_name), column_sql)
        ) as copy:
            with path.open("rb") as source:
                while chunk := source.read(1024 * 1024):
                    copy.write(chunk)

        cursor.execute(sql.SQL("SELECT count(*) FROM {}").format(sql.Identifier(temp_name)))
        received = cursor.fetchone()[0]
        # EXCEPT сравнивает два набора сразу. Это важно для больших выгрузок:
        # предыдущий вариант искал каждую строку staging во всей таблице.
        cursor.execute(
            sql.SQL(
                "INSERT INTO {} ({}) "
                "SELECT DISTINCT {} FROM {} "
                "EXCEPT SELECT {} FROM {}"
            ).format(
                target, column_sql, column_sql, sql.Identifier(temp_name),
                column_sql, target,
            )
        )
        inserted = cursor.rowcount
    return {
        "staging_file": path.name,
        "table": f"{schema}.{table}",
        "received_rows": received,
        "inserted_rows": inserted,
        "skipped_existing_rows": received - inserted,
    }


def load_staging_to_database(staging_dir: Path, database_url: str) -> list[dict]:
    """Переносит staging-файлы текущей загрузки в PostgreSQL без дублей."""
    files = [
        path for path in sorted(staging_dir.glob("*.csv"))
        if path.name in STAGING_TABLES
    ]
    if not files:
        raise ValueError("В этой загрузке нет staging CSV для записи в базу.")

    results = []
    with psycopg.connect(database_url) as connection:
        for path in files:
            schema, table = STAGING_TABLES[path.name]
            with connection.cursor() as cursor:
                cursor.execute(
                    sql.SQL("SELECT count(*) FROM {}").format(
                        sql.Identifier(schema, table)
                    )
                )
                existing_rows = cursor.fetchone()[0]
            if existing_rows and table != "boarding_passes":
                raise ValueError(
                    f"{schema}.{table} уже содержит "
                    f"{existing_rows} строк. "
                    "Повторная загрузка в эту таблицу запрещена."
                )

            results.append(
                _load_one_staging_file(
                    connection,
                    path,
                    schema,
                    table,
                )
            )
    return results

def _stage_boarding_passes_batch(
    source_files: list[Path],
    batch_number: int,
    job_dir: Path,
    parsed_root: Path,
    staging_root: Path,
) -> list[dict]:
    """
    Запускает парсер для группы XLSX-файлов и объединяет CSV.
    Исходные файлы не изменяются.
    """
    if not source_files:
        return []

    script = PARSERS_DIR / "parse_boarding_passes.py"

    if not script.is_file():
        raise FileNotFoundError(
            f"Парсер не найден: {script}"
        )

    input_dir = (
        job_dir / "_batch_input" / f"batch_{batch_number:03d}"
    )
    output_dir = (
        parsed_root / f"boarding_passes_batch_{batch_number:03d}"
    )

    shutil.rmtree(input_dir, ignore_errors=True)
    shutil.rmtree(output_dir, ignore_errors=True)

    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    items = []

    try:
        # Создаём ссылки на исходные файлы, чтобы не копировать
        # сотни больших XLSX без необходимости.
        for source in source_files:
            target = input_dir / source.name

            try:
                os.link(source, target)
            except OSError:
                shutil.copy2(source, target)

        command = [
            sys.executable,
            "-X", "utf8",
            str(script),
            "--input", str(input_dir),
            "--output", str(output_dir),
            "--overwrite",
        ]

        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"

        result = subprocess.run(
            command,
            cwd=PROJECT_DIR,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=21600,
            check=False,
            env=env,
        )

        parse_report_path = output_dir / "parse_report.json"

        if not parse_report_path.is_file():
            raise RuntimeError(
                "Парсер не создал parse_report.json. "
                f"Код завершения: {result.returncode}. "
                f"Лог: {(result.stdout + result.stderr)[-3000:]}"
            )

        parse_records = json.loads(
            parse_report_path.read_text(encoding="utf-8")
        )
        records_by_name = {
            row.get("source_file"): row
            for row in parse_records
            if row.get("source_file")
        }

        staging_target = staging_root / "boarding_passes.csv"

        for source in source_files:
            parsed_csv = output_dir / f"{source.stem}_table.csv"
            parsed_record = records_by_name.get(source.name)

            message = None
            staged = False

            if parsed_csv.is_file():
                try:
                    _append_csv(
                        parsed_csv,
                        staging_target,
                        delimiter=";",
                    )
                    staged = True
                except Exception as exc:
                    message = str(exc)
            else:
                message = "Парсер не создал CSV для этого файла."

            if parsed_record is None:
                staged = False
                message = message or (
                    "Файл отсутствует в parse_report.json."
                )
            elif parsed_record.get("file_failed"):
                staged = False
                message = parsed_record["file_failed"]
            elif parsed_record.get("sheets_failed", 0) > 0:
                staged = False
                message = (
                    "Не все листы обработаны. "
                    f"Ошибок: {parsed_record['sheets_failed']}."
                )

            item = {
                "source": str(source.relative_to(job_dir)).replace(
                    "\\", "/"
                ),
                "extension": ".xlsx",
                "status": "staged" if staged else "error",
                "output_files": (
                    [
                        str(parsed_csv.relative_to(job_dir)).replace(
                            "\\", "/"
                        )
                    ]
                    if parsed_csv.is_file()
                    else []
                ),
                "staging_files": (
                    ["staging/boarding_passes.csv"]
                    if parsed_csv.is_file()
                    else []
                ),
                "rows_parsed": (
                    parsed_record.get("rows_parsed", 0)
                    if parsed_record
                    else 0
                ),
                "rows_with_warnings": (
                    parsed_record.get("rows_with_warnings", 0)
                    if parsed_record
                    else 0
                ),
                "message": message or "Файл обработан.",
            }
            items.append(item)

    finally:
        # Удаляем только временные ссылки или копии входных XLSX.
        shutil.rmtree(input_dir, ignore_errors=True)

    return items


def run_import(job_dir: Path) -> dict:
    """Автоматически обрабатывает прямые файлы и ZIP-архивы."""

    parsed_root = job_dir / "parsed"
    staging_root = job_dir / "staging"

    parsed_root.mkdir(parents=True, exist_ok=True)
    staging_root.mkdir(parents=True, exist_ok=True)

    archives = [
        path for path in job_dir.iterdir()
        if path.is_file() and path.suffix.lower() == ".zip"
    ]

    extracted_roots = []
    archive_reports = []

    for index, archive_path in enumerate(sorted(archives), start=1):
        destination = (
            job_dir / "extracted"
            / f"{archive_path.stem}_{index:03d}"
        )

        shutil.rmtree(destination, ignore_errors=True)
        extracted = _extract_zip_safely(
            archive_path,
            destination,
        )

        extracted_roots.append(destination)
        archive_reports.append({
            "archive": archive_path.name,
            "extracted_files": len(extracted),
        })

    sources = [
        path for path in job_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() != ".zip"
        and path.name != "import_report.json"
    ]

    for root in extracted_roots:
        sources.extend(
            path for path in root.rglob("*")
            if path.is_file()
        )

    supported_extensions = set(PARSER_SCRIPTS) | {".csv"}

    ignored_files = [
        str(path.relative_to(job_dir)).replace("\\", "/")
        for path in sources
        if path.suffix.lower() not in supported_extensions
    ]

    sources = sorted(
        path for path in sources
        if path.suffix.lower() in supported_extensions
    )

    report = {
        "job_id": job_dir.name,
        "status": "completed",
        "started_at": datetime.now(timezone.utc).isoformat(),
        "archives": archive_reports,
        "ignored_files": ignored_files,
        "files": [],
    }

    xlsx_groups: dict[Path, list[Path]] = {}

    # Сначала обычные файлы и Excel-выгрузки BoardingData.
    for source in sources:
        extension = source.suffix.lower()

        item = {
            "source": str(source.relative_to(job_dir)).replace(
                "\\", "/"
            ),
            "extension": extension,
        }

        if extension == ".xlsx":
            staged = _stage_boarding_data_xlsx(
                source,
                staging_root,
            )

            if staged:
                item.update(
                    status="staged",
                    staging_files=staged,
                    message="BoardingData подготовлен для PostgreSQL.",
                )
                report["files"].append(item)
            else:
                xlsx_groups.setdefault(
                    source.parent, []
                ).append(source)

            continue

        if extension == ".csv":
            staged = _stage_boarding_data_csv(
                source,
                staging_root,
            )

            if staged:
                item.update(
                    status="staged",
                    staging_files=staged,
                    message="BoardingData CSV подготовлен.",
                )
            else:
                item.update(
                    status="saved_without_parsing",
                    message=(
                        "CSV сохранён, но его структура не распознана "
                        "как BoardingData."
                    ),
                )

            report["files"].append(item)
            continue

        try:
            output = parsed_root / source.stem
            shutil.rmtree(output, ignore_errors=True)
            output.mkdir(parents=True, exist_ok=True)

            result = subprocess.run(
                _parser_command(source, output),
                cwd=PROJECT_DIR,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=21600,
                check=False,
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            )

            item["log"] = (
                result.stdout + result.stderr
            )[-8000:]

            item["output_files"] = [
                str(path.relative_to(job_dir)).replace("\\", "/")
                for path in output.rglob("*")
                if path.is_file()
            ]

            if result.returncode != 0:
                item.update(
                    status="error",
                    message="Парсер завершился с ошибкой.",
                )
            else:
                staged = _stage_parsed_files(
                    source,
                    output,
                    staging_root,
                )
                item.update(
                    status="staged",
                    staging_files=staged,
                )

            report["files"].append(item)

        except Exception as exc:
            item.update(
                status="error",
                message=str(exc),
            )
            report["files"].append(item)

    # Теперь пакетно обрабатываем все посадочные талоны.
    for batch_number, (parent, group) in enumerate(
        sorted(xlsx_groups.items(), key=lambda entry: str(entry[0])),
        start=1,
    ):
        try:
            items = _stage_boarding_passes_batch(
                group,
                batch_number,
                job_dir,
                parsed_root,
                staging_root,
            )
            report["files"].extend(items)

        except Exception as exc:
            for source in group:
                report["files"].append({
                    "source": str(source.relative_to(job_dir)).replace(
                        "\\", "/"
                    ),
                    "extension": ".xlsx",
                    "status": "error",
                    "message": str(exc),
                })

    statuses = [
        item.get("status")
        for item in report["files"]
    ]

    if statuses and all(status == "staged" for status in statuses):
        report["status"] = "completed"
    elif statuses and all(status == "error" for status in statuses):
        report["status"] = "failed"
    else:
        report["status"] = "partial"

    report["finished_at"] = datetime.now(timezone.utc).isoformat()

    (job_dir / "import_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return report
