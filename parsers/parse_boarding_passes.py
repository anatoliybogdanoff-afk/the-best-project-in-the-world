
import argparse
import csv
import json
import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = {
    "document_title": "A1",
    "sequence_number": "H1",
    "title": "A3",
    "passenger_name": "B3",
    "class_code": "H3",
    "flight_number": "A5",
    "origin_city": "D5",
    "destination_city": "H5",
    "gate": "B7",
    "origin_airport": "D7",
    "destination_airport": "H7",
    "flight_date": "A9",
    "flight_time": "C9",
    "operator_note": "E9",
    "boarding_note": "A11",
    "seat": "H11",
    "booking_code": "B13",
    "ticket_number": "E13",
}
META = ["source_file", "source_sheet", "file_date", "parse_warnings", "raw_cells_json", "raw_types_json"]
HEADERS = list(FIELDS) + META
EXPECTED_LABELS = {"A1": "BOARDING PASS", "F1": "SEQUENCE:", "A7": "GATE", "G11": "SEAT", "A13": "PNR:", "D13": "E-TICKET"}
REQUIRED = ["passenger_name", "flight_number", "origin_airport", "destination_airport", "flight_date", "flight_time", "booking_code", "ticket_number"]


def source_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_sheet(sheet, source_file):

    sheet.reset_dimensions()
    values, types = {}, {}
    for cells in sheet.iter_rows():
        for cell in cells:
            if cell.value is not None:
                values[cell.coordinate] = source_text(cell.value)
                types[cell.coordinate] = cell.data_type
    row = {field: values.get(address, "") for field, address in FIELDS.items()}
    warnings = []
    for address, expected in EXPECTED_LABELS.items():
        if values.get(address, "").strip() != expected:
            warnings.append(f"Unexpected template label at {address}")
    for field in REQUIRED:
        if row[field].strip().upper() in {"", "N/A"}:
            warnings.append(f"Missing field: {field}")
    match = re.search(r"(\d{4}-\d{2}-\d{2})", source_file.name)
    file_date = match.group(1) if match else ""
    if file_date and row["flight_date"][:10] != file_date:
        warnings.append("Date in pass differs from filename")
    if types.get("E13") == "n":
        warnings.append("Ticket stored as number in source; original leading zeros/precision cannot be guaranteed")
    if any(types.get(address) == "f" for address in FIELDS.values()):
        warnings.append("Formula in a mapped field: preserved as text, not calculated")
    row.update(source_file=source_file.name, source_sheet=sheet.title, file_date=file_date,
               parse_warnings="; ".join(warnings),
               raw_cells_json=json.dumps(values, ensure_ascii=False),
               raw_types_json=json.dumps(types, ensure_ascii=False))
    if any(len(v) > 32767 for v in row.values()):
        raise ValueError("A field exceeds Excel's 32767-character cell limit")
    return row


def append_text(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.data_type = "s"
        cell.number_format = "@"


def style_table(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25476A")
    for index, header in enumerate(next(sheet.values), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(36, max(18, len(str(header)) + 2))


def process_file(source, target, overwrite=False, check_only=False):
    xlsx = target / (source.stem + "_table.xlsx")
    csv_path = target / (source.stem + "_table.csv")
    if not check_only and not overwrite and (xlsx.exists() or csv_path.exists()):
        raise FileExistsError(f"Output exists for {source.name}; use --overwrite intentionally")
    src = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    rows, errors = [], []
    try:
        total = len(src.worksheets)
        for sheet in src.worksheets:
            try:
                rows.append(parse_sheet(sheet, source))
            except Exception as exc:
                errors.append([source.name, sheet.title, f"{type(exc).__name__}: {exc}"])
    finally:
        src.close()
    assert total == len(rows) + len(errors)
    result = {"source_file": source.name, "sheets_read": total, "rows_parsed": len(rows),
              "rows_with_warnings": sum(bool(r["parse_warnings"]) for r in rows), "sheets_failed": len(errors)}
    if check_only:
        return result, errors
    if len(rows) > 1048575:
        raise ValueError("Too many rows for one Excel sheet; split this source before export")
    out = Workbook()
    data = out.active
    data.title = "Passes"
    append_text(data, HEADERS)
    for row in rows:
        append_text(data, [row[h] for h in HEADERS])
    err = out.create_sheet("Errors")
    append_text(err, ["source_file", "source_sheet", "error"])
    for error in errors:
        append_text(err, error)
    check = out.create_sheet("Check")
    append_text(check, ["metric", "value"])
    for key, value in result.items():
        append_text(check, [key, str(value)])
    for sheet in out:
        style_table(sheet)
    temp = xlsx.with_suffix(".tmp.xlsx")
    out.save(temp)
    out.close()
    temp.replace(xlsx)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=HEADERS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return result, errors


def main():
    parser = argparse.ArgumentParser(description="Flatten daily boarding-pass XLSX books into daily tables")
    parser.add_argument("--input", required=True, type=Path, help="Folder containing original daily XLSX files")
    parser.add_argument("--output", required=True, type=Path, help="A DIFFERENT folder for generated tables")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing previously generated outputs")
    parser.add_argument("--check-only", action="store_true", help="Read and validate without writing workbooks")
    args = parser.parse_args()
    input_dir, output_dir = args.input.resolve(), args.output.resolve()
    if not input_dir.is_dir():
        parser.error(f"Input folder not found: {input_dir}")
    if input_dir == output_dir:
        parser.error("Input and output folders must differ")

    files = sorted(p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$") and not p.stem.endswith("_table"))
    if not files:
        parser.error("No XLSX files found in input folder (not recursive)")
    if not args.check_only:
        output_dir.mkdir(parents=True, exist_ok=True)
        if not args.overwrite and any((output_dir / (p.stem + suffix)).exists() for p in files for suffix in ("_table.xlsx", "_table.csv")):
            parser.error("Some outputs already exist; use another folder or --overwrite")
    report, all_errors = [], []
    for index, source in enumerate(files, 1):
        try:
            result, errors = process_file(source, output_dir, args.overwrite, args.check_only)
            report.append(result)
            all_errors.extend(errors)
            print(f"[{index}/{len(files)}] {source.name}: {result['rows_parsed']} rows, {result['sheets_failed']} errors", flush=True)
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            report.append({"source_file": source.name, "file_failed": message})
            all_errors.append([source.name, "", message])
            print(f"ERROR: {source.name}: {message}", flush=True)
    if not args.check_only:
        (output_dir / "parse_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        with (output_dir / "errors.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(["source_file", "source_sheet", "error"])
            writer.writerows(all_errors)
    print("Finished. Check parse_warnings, Errors and parse_report.json.")
    if all_errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
